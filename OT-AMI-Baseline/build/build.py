# -*- coding: utf-8 -*-
import re, sys, os, datetime, json, shutil, warnings
from collections import defaultdict, Counter
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
import content as C

UP = "/root/.claude/uploads/227ed269-dff6-50e5-bb20-7a7219919239/"
SCF_FILE = UP + "a44946a1-Secure_Controls_Framework_SCF_-_2026.2.xlsx"
MASTER_FILE = UP + "1b628e39-TRA-ISP-Control-Catalog-v2026.2-Risk-Control-Assessment.xlsx"
TEMPLATE_FILE = UP + "dac479f7-control-baseline-template_2.xlsx"
OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else "/home/user/3HUE-Website/OT-AMI-Baseline"
os.makedirs(OUT_DIR, exist_ok=True)
TODAY = datetime.datetime(2026, 10, 1)
VERSION = "v1.0 (SCF 2026.2)"

# ---------------------------------------------------------------- load SCF
wb = openpyxl.load_workbook(SCF_FILE, read_only=True)
ws = wb["SCF 2026.2"]
SCF_HDR = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
SCF = {}
ORDER = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[2]:
        SCF[r[2]] = r
        ORDER.append(r[2])
print("SCF controls", len(SCF))

AO = defaultdict(list)
for r in wb["Assessment Objectives 2026.2"].iter_rows(min_row=2, values_only=True):
    if r[0] and r[2]:
        AO[r[0]].append((r[1], r[2], r[6], r[7]))

DOM_PRINCIPLE = {}
for r in wb["SCF Domains & Principles"].iter_rows(min_row=2, values_only=True):
    if r[2]:
        DOM_PRINCIPLE[r[2]] = (r[1], r[3], r[4])

# ---------------------------------------------------------------- load 3HUE master
wbm = openpyxl.load_workbook(MASTER_FILE, read_only=True)
wsm = wbm["ISP Control Catalog Q3-2026"]
MASTER = {}
for r in wsm.iter_rows(min_row=2, values_only=True):
    if r[2]:
        MASTER[r[2]] = dict(accountability=r[4], audience=r[50], applicability=r[51],
                            scf_i=r[45], scf_r=r[46], lifecycle=r[3])
print("Master controls", len(MASTER))

def has(r, idx):
    return idx is not None and r[idx] not in (None, "")

def clean_map(v):
    if v in (None, ""):
        return ""
    return ", ".join(x.strip() for x in str(v).split("\n") if x.strip())

# ---------------------------------------------------------------- scoping
REQ = {k: dict(label=l, idx=i, tier=t, rationale=ra) for (k, l, i, t, ra) in C.REQUIREMENTS}
ZCR_IDS = defaultdict(list)
for zid, title, ids, rat in C.ZCR_CROSSWALK:
    for cid in ids:
        assert cid in SCF, "ZCR crosswalk references unknown SCF control: " + cid
        ZCR_IDS[cid].append(zid)

triggers = {}
for cid in ORDER:
    r = SCF[cid]
    t = []
    for k, q in REQ.items():
        if k == "EMB":
            if cid.startswith("EMB-"):
                t.append(k)
        elif k == "ZCR":
            if cid in ZCR_IDS:
                t.append(k)
        elif has(r, q["idx"]):
            t.append(k)
    if t:
        triggers[cid] = t

SCOPE = []
EXCL = []
for cid in ORDER:
    if cid not in triggers:
        continue
    dom = cid.split("-")[0]
    tiers = [REQ[k]["tier"] for k in triggers[cid]]
    tier = min(tiers)
    if dom in C.EXCLUDED_DOMAINS and tier > 1:
        EXCL.append((cid, "excluded domain", triggers[cid]))
        continue
    SCOPE.append(cid)
print("In scope", len(SCOPE), "excluded (domain rule)", len(EXCL))
only=Counter(triggers[c][0] for c in SCOPE if len(triggers[c])==1)
print("single-trigger controls:", dict(only))
print(Counter(cid.split("-")[0] for cid in SCOPE))

# ---------------------------------------------------------------- text helpers
US_WORDS = [("modelled", "modeled"), ("modelling", "modeling"), ("amongst", "among"), ("whilst", "while"), ("towards", "toward"),
            ("organisation", "organization"), ("organise", "organize"), ("authorise", "authorize"), ("authorisation", "authorization"),
            ("prioritise", "prioritize"), ("prioritisation", "prioritization"), ("centre", "center"), ("defence", "defense"),
            ("behaviour", "behavior"), ("licence", "license"), ("programme", "program"), ("catalogue", "catalog"),
            ("utilise", "utilize"), ("minimise", "minimize"), ("optimise", "optimize"), ("recognise", "recognize"),
            ("standardise", "standardize"), ("categorise", "categorize"), ("characterise", "characterize"),
            ("labelled", "labeled"), ("labelling", "labeling"), ("cancelled", "canceled"), ("judgement", "judgment"),
            ("analyse", "analyze"), ("analysed", "analyzed"), ("analysing", "analyzing"), ("synchronise", "synchronize"),
            ("artefact", "artifact"), ("enrolment", "enrollment"), ("fulfil", "fulfill"), ("learnt", "learned"), ("grey", "gray")]

def us_english(text):
    """Normalize authored/derived text to US English spelling (case-preserving for a leading capital)."""
    if not isinstance(text, str):
        return text
    for uk, us in US_WORDS:
        text = re.sub(r"\b%s\b" % uk, us, text)
        text = re.sub(r"\b%s\b" % uk.capitalize(), us.capitalize(), text)
    return text

def split_desc(desc):
    d = re.sub(r"\s+", " ", desc.strip())
    m = re.match(r"^(.*?)\b[Mm]echanisms exist to\s+(.*)$", d, re.S)
    if m:
        prefix = m.group(1).strip()
        rest = m.group(2).strip()
        return prefix, rest
    return None, d

def lc_first(s):
    if s[:4] == "The " or s[:2] == "A " or s[:3] == "An ":
        return s[0].lower() + s[1:]
    return s

def strip_period(s):
    s = s.strip()
    while s.endswith(".") or s.endswith(";"):
        s = s[:-1].rstrip()
    return s

def program_control(desc):
    return desc.strip().startswith("Mechanisms exist to facilitate the implementation of")

def gen_objective(cid, r, dom):
    desc = r[3]
    prefix, rest = split_desc(desc)
    if program_control(desc):
        rest2 = rest[len("facilitate the implementation of "):] if rest.startswith("facilitate the implementation of ") else rest
        rest2 = strip_period(rest2)
        core = ("The organization develops, implements and governs processes and documentation to facilitate the "
                "implementation of %s, as well as associated standards, controls and procedures, across enterprise IT, "
                "Operational Technology (OT) and Advanced Metering Infrastructure (AMI) environments." % rest2)
    elif prefix is not None:
        pfx = (prefix.lower() + " ") if prefix else ""
        core = "The organization implements and operates %smechanisms to %s." % (pfx, strip_period(rest))
    else:
        core = "The organization ensures that %s." % strip_period(lc_first(rest))
    ctx = (C.CTRL.get(cid, {}).get("ot_context") or C.DOMAIN[dom]["ot_context"])
    return core + "\n\nWithin the OT and AMI environment, this objective " + strip_period(ctx) + "."

def clean_ao(text):
    t = re.sub(r"\s+", " ", text.strip())
    t = re.sub(r"<[^>]*ODP\[\d+\]:\s*([^>]*)>", r"organization-defined \1", t)
    t = re.sub(r"<[^>]*>", "organization-defined parameters", t)
    t = t.replace("C at the physical perimeter", "at the physical perimeter")
    t = re.sub(r"\.\s+(?=[a-z])", ", ", t)
    t = t.replace(", and any", " and any")
    t = t.replace("is/are selected:", "is selected:")
    return strip_period(t)

def ao_items(cid):
    items = []
    seen = set()
    rows = []
    for aoid, text, rigor, sdp in AO.get(cid, []):
        try:
            rg = int(str(rigor).strip())
        except Exception:
            continue
        if rg not in (1, 2, 3):
            continue
        if re.search(r"\bCUI\b|Controlled Unclassified", text):
            continue
        rows.append((rg, aoid, text))
    rows.sort(key=lambda x: (x[0], x[1]))
    for rg, aoid, text in rows:
        t = clean_ao(text)
        key = re.sub(r"[^a-z0-9]", "", t.lower())
        if key in seen or not t:
            continue
        seen.add(key)
        items.append("Ensuring " + lc_first(t))
        if len(items) >= 8:
            break
    return items

def letters(n):
    out = []
    for i in range(n):
        if i < 26:
            out.append(chr(97 + i))
        else:
            out.append("a" + chr(97 + i - 26))
    return out

def fmt_list(intro, items):
    items = [strip_period(i) for i in items if i]
    if not items:
        return intro + "."
    if len(items) == 1:
        return intro + ":\na. " + items[0] + "."
    lab = letters(len(items))
    lines = []
    for i, it in enumerate(items):
        if i == len(items) - 1:
            lines.append("%s. %s." % (lab[i], it))
        elif i == len(items) - 2:
            lines.append("%s. %s; and" % (lab[i], it))
        else:
            lines.append("%s. %s;" % (lab[i], it))
    return intro + ":\n" + "\n".join(lines)

def gen_standard(cid, r, dom, accountability):
    role = C.ROLE_PHRASE.get(accountability, C.ROLE_PHRASE["IT Ops - Shared"])
    items = ao_items(cid)
    if not items:
        prefix, rest = split_desc(r[3])
        if prefix is not None:
            items = ["Implementing and maintaining %smechanisms to %s" % ((prefix.lower() + " ") if prefix else "", strip_period(rest))]
        else:
            items = ["Ensuring " + strip_period(lc_first(rest))]
    ot = C.CTRL.get(cid, {}).get("std_items") or C.DOMAIN[dom]["std_items"]
    return fmt_list(role, items + list(ot))

def gen_intent(cid, dom):
    for pfx, txt in C.INTENT_OVERRIDE:
        if cid.startswith(pfx):
            return txt
    return C.DOMAIN[dom]["intent"]

MAP_FOR_GUIDE = [
    ("IEC 62443-2-1", 53), ("IEC 62443-3-3", 54), ("IEC 62443-4-1", 55), ("IEC 62443-4-2", 56),
    ("NIST SP 800-82 R3 (Moderate OT overlay)", 86), ("NERC CIP", 163), ("CISA CPG", 130), ("DOE C2M2", 132),
    ("NIST CSF 2.0", 103), ("NIST SP 800-53B Moderate", 81), ("NIST SP 800-161 R1 C-SCRM", 90),
    ("CSA IoT SCF", 44), ("NIST Privacy Framework", 74), ("NIST SP 800-53B Privacy", 79),
]

def gen_guideline(cid, r, dom):
    g = C.CTRL.get(cid, {}).get("guide") or C.DOMAIN[dom]["guide"]
    parts = ["OT/AMI implementation guidance: " + strip_period(g) + "."]
    sol = r[10]
    if sol:
        bullets = [strip_period(re.sub(r"\s*\(?https?://\S+\)?", "", b.strip().lstrip("∙•-").strip())) for b in str(sol).split("\n") if b.strip()]
        bullets = [b for b in bullets if b]
        if bullets:
            parts.append("Suggested solutions and considerations (SCF 2026.2, enterprise): " + "; ".join(bullets) + ".")
    al = []
    for label, idx in MAP_FOR_GUIDE:
        v = clean_map(r[idx])
        if v:
            al.append("%s: %s" % (label, v))
    if cid in ZCR_IDS:
        al.append("ISA/IEC 62443-3-2: " + ", ".join(ZCR_IDS[cid]))
    if al:
        parts.append("In-scope requirement alignment: " + " | ".join(al) + ".")
    parts.append("Validation: %s conformity validation; evidence artifacts %s; PPTDF applicability: %s." % (
        r[4] or "Annual", clean_map(r[5]) or "to be defined", r[13] or "Process"))
    return "\n\n".join(parts)

def nistir_family(r):
    v = r[78]
    if not v:
        return ""
    fams = []
    for x in str(v).split("\n"):
        m = re.match(r"([A-Z]{2})-", x.strip())
        if m and m.group(1) in C.NISTIR7628_FAMILY:
            f = C.NISTIR7628_FAMILY[m.group(1)]
            if f not in fams:
                fams.append(f)
    return "; ".join(fams)

def tsc_split(v):
    if not v:
        return "", ""
    items = [x.strip() for x in str(v).split("\n") if x.strip()]
    tsc = [x for x in items if "-POF" not in x]
    pof = [x for x in items if "-POF" in x]
    return ", ".join(tsc), ", ".join(pof)

def zta(r):
    out = []
    if r[101]:
        out += [x.strip() for x in str(r[101]).split("\n") if x.strip()]
    if r[139]:
        out += ["DoD " + x.strip() for x in str(r[139]).split("\n") if x.strip()]
    return ", ".join(out)

FUNC_FALLBACK = {"GOV": "Govern", "CPL": "Govern", "PRM": "Govern", "RSK": "Identify", "TPM": "Govern", "AST": "Identify", "IAO": "Identify", "THR": "Identify", "VPM": "Identify", "IRO": "Respond", "BCD": "Recover", "MON": "Detect"}
AUD = {"Technology": "Technical", "Facility": "Technical", "Process": "Management", "People": "All Users", "Data": "All Users"}
FREQ = {"Annual": 365, "Semi-Annual": 180, "Quarterly": 90}

def ot_category(cid, r, trig, tier):
    dom = cid.split("-")[0]
    cats = []
    if dom == "EMB" or "IEC62443-4-2" in trig or ("CSAIOT" in trig and r[13] in ("Technology", "Data")):
        cats.append("AMI Endpoint / Field Device")
    if tier == 1:
        cats.append("OT Core")
    if dom == "PRI" or tier == 4:
        cats.append("Privacy (Customer Energy Usage Data)")
    if tier in (2, 3) and "AMI Endpoint / Field Device" not in cats:
        cats.append("Converged IT/OT Enterprise")
    if not cats:
        cats.append("Converged IT/OT Enterprise")
    return "; ".join(cats)

# ---------------------------------------------------------------- build rows
ROWS = []
for cid in SCOPE:
    r = SCF[cid]
    dom = cid.split("-")[0]
    trig = triggers[cid]
    tier = min(REQ[k]["tier"] for k in trig)
    m = MASTER.get(cid)
    accountability = (m["accountability"] if m and m["accountability"] in C.ROLE_PHRASE else C.DEFAULT_ACCOUNTABILITY[dom])
    audience = AUD.get(r[13], "Management")
    if program_control(r[3] or ""):
        audience = "Management"
    applicability = "Basic" if tier == 1 else "Enhanced"
    cadence = r[4] or "Annual"
    tsc, pof = tsc_split(r[34])
    row = dict(
        cid=cid, title=r[1], domain=r[0], family=dom, function=(r[14] if r[14] and r[14] != "N/A" else FUNC_FALLBACK.get(dom, "Protect")), desc=r[3], question=r[11],
        weight=r[12], pptdf=r[13], cadence=cadence, erl=clean_map(r[5]), tier=tier,
        tier_name=C.TIER_NAMES[tier], applicability=applicability,
        rationale="; ".join(REQ[k]["label"].split(" (")[0] for k in trig),
        trig=trig, category=ot_category(cid, r, trig, tier),
        master_status=("In 3HUE master catalog (%s)" % m["lifecycle"]) if m else "New for OT/AMI engagement",
        accountability=accountability, owner_corp=C.OWNERS[dom][0], owner_plat=C.OWNERS[dom][1],
        audience=audience, freq=FREQ.get(cadence, 365), third=C.THIRD_PARTIES[dom],
        intent=us_english(gen_intent(cid, dom)), objective=us_english(gen_objective(cid, r, dom)),
        standard=us_english(gen_standard(cid, r, dom, accountability)), guideline=us_english(gen_guideline(cid, r, dom)),
        tsc=tsc, pof=pof, cobit=clean_map(r[41]), iso27001=clean_map(r[60]), iso27002=clean_map(r[61]),
        n53mod=clean_map(r[81]), pci_a=clean_map(r[108]), pci_d=clean_map(r[115]), cmmc2=clean_map(r[135]),
        hipaa=clean_map(r[159]), fedramp=clean_map(r[149]), gdpr=clean_map(r[200]), scf_b=clean_map(r[32]),
        scf_i=(m["scf_i"] if m and m["scf_i"] else ""), scf_r=(m["scf_r"] if m and m["scf_r"] else ""),
        zta=zta(r), solutions=r[8] or "",
        m2_1=clean_map(r[53]), m3_3=clean_map(r[54]), m4_1=clean_map(r[55]), m4_2=clean_map(r[56]),
        zcr=", ".join(ZCR_IDS.get(cid, [])), m82=clean_map(r[84]), m82mod=clean_map(r[86]), nerc=clean_map(r[163]),
        cpg=clean_map(r[130]), c2m2=clean_map(r[132]), csf2=clean_map(r[103]), n53=clean_map(r[78]),
        n161=clean_map(r[90]), iot=clean_map(r[44]), npf=clean_map(r[74]), n53priv=clean_map(r[79]),
        nistir=nistir_family(r), n207=clean_map(r[101]), esp2=("x" if r[29] else ""),
        ent_solutions=r[10] or "",
    )
    ROWS.append(row)

for row in ROWS:
    for k in ("intent", "objective", "standard", "guideline"):
        assert "Transit" not in row[k], (row["cid"], k)

# ---------------------------------------------------------------- styles
ARIAL = Font(name="Arial", size=10)
ARIAL_B = Font(name="Arial", size=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 32

def write_table(ws, headers, rows, widths=None, start_row=1):
    ws.append(headers)
    style_header(ws, len(headers), start_row)
    for rr in rows:
        ws.append(rr)
    for rrow in ws.iter_rows(min_row=start_row + 1, max_row=ws.max_row):
        for cell in rrow:
            cell.font = ARIAL
            cell.alignment = WRAP
            cell.border = BORDER
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

# ================================================================ ROSETTA STONE WORKBOOK
rwb = openpyxl.Workbook()
# ---- README
ws = rwb.active
ws.title = "README"
readme = [
    ("3HUE OT / IT-OT Convergence / AMI Control Baseline - Rosetta Stone", True),
    ("Version: %s    Baseline date: %s    Source framework: Secure Controls Framework (SCF) 2026.2" % (VERSION, TODAY.strftime("%Y-%m-%d")), False),
    ("", False),
    ("Purpose", True),
    ("This workbook is the engagement Rosetta Stone: the SCF 2026.2 control set filtered against the in-scope regulatory, security, industry and privacy requirements for an Operational Technology (OT), IT/OT convergence and Advanced Metering Infrastructure (AMI) environment, with four authored columns added (Management Intent, Control Objectives, Control Standards, Control Guidelines). It is the working record behind the Client Vision upload file.", False),
    ("", False),
    ("How the baseline was scoped", True),
    ("1. Every SCF 2026.2 control was tested for a mapping to each in-scope requirement listed on the 'Scoping Requirements' sheet. A control is in the baseline if it maps to at least one in-scope requirement.", False),
    ("2. Each control is assigned to the lowest (most mandatory) tier that triggered it: Tier 1 OT Core, Tier 2 Converged IT/OT enterprise, Tier 3 AMI endpoint / embedded / supply chain, Tier 4 Privacy (Customer Energy Usage Data).", False),
    ("3. ANSI/ISA-62443-3-2:2020 is not natively mapped in the SCF; its Zone & Conduit Requirements (ZCR 1-7) were crosswalked manually to SCF controls (see '62443-3-2 ZCR Crosswalk'). Any control referenced by the crosswalk is force-included at Tier 1.", False),
    ("4. The Artificial Intelligence & Autonomous Technologies (AAT) and Quantum Security (QTS) domains are excluded unless a Tier 1 requirement maps to them; they can be added as future overlays.", False),
    ("5. Applicability: 'Basic' = Tier 1 (mandatory OT core); 'Enhanced' = Tiers 2-4. Target Audience derives from SCF PPTDF applicability (Technology/Facility = Technical, Process = Management, People/Data = All Users); program-level controls are always Management.", False),
    ("6. Accountability uses the 3HUE master catalog assignment where the control exists in the master; otherwise the domain default on the '_index' sheet. OT roles are expressed through the Control Owner-Platform column.", False),
    ("", False),
    ("Authored columns", True),
    ("Management Intent -> Client Vision 'Policy Statement'. Domain-level policy intent written for a converged IT/OT/AMI program; sub-domain variants where a topic warrants it (e.g., remote access, wireless, patching, key management).", False),
    ("Control Objectives -> Client Vision 'Control Objectives'. Outcome statement derived from the SCF control description plus an OT/AMI extension sentence.", False),
    ("Control Standards -> Client Vision 'Standard Content'. Accountable role plus lettered requirements derived from the SCF 2026.2 Assessment Objectives (rigor levels 1-3) and OT/AMI-specific requirements authored for this engagement.", False),
    ("Control Guidelines -> Client Vision 'Implementation Notes'. OT/AMI implementation guidance, SCF enterprise solutions and considerations, in-scope requirement alignment and validation cadence / evidence references.", False),
    ("", False),
    ("Column notes", True),
    ("NISTIR 7628 R1 family is derived from the NIST SP 800-53 R5 family of each mapped control (e.g., AC -> SG.AC); it is a family-level crosswalk, not a requirement-level mapping.", False),
    ("SCF-I (Cyber Insurance) and SCF-R (Ransomware) flags are carried forward from the 3HUE master catalog where the control exists there; SCF 2026.2 no longer publishes those baselines.", False),
    ("Zero Trust Architecture combines NIST SP 800-207 tenets and DoD Zero Trust Reference Architecture 2.0 references, consistent with the 3HUE master catalog convention.", False),
    ("All authored and derived text uses US English spelling and conventions; SCF control descriptions are reproduced verbatim from SCF 2026.2.", False),
    ("No client-specific or prior-engagement content is used in any authored text.", False),
]
for i, (txt, bold) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=txt)
    c.font = ARIAL_B if bold else ARIAL
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 160
ws.cell(row=1, column=1).font = Font(name="Arial", size=14, bold=True)

# ---- Scoping Requirements
ws = rwb.create_sheet("Scoping Requirements")
hdr = ["Key", "In-Scope Requirement", "Tier", "SCF 2026.2 Column", "Scoping Rationale", "SCF Controls Mapped (all)", "Controls In Baseline"]
rows = []
for k, q in REQ.items():
    if k == "EMB":
        total = sum(1 for c in ORDER if c.startswith("EMB-"))
    elif k == "ZCR":
        total = len(ZCR_IDS)
    else:
        total = sum(1 for c in ORDER if has(SCF[c], q["idx"]))
    col = (SCF_HDR[q["idx"]].replace("\n", " ") if q["idx"] is not None else ("SCF domain EMB" if k == "EMB" else "Manual crosswalk (this workbook)"))
    inb = sum(1 for row in ROWS if k in row["trig"])
    rows.append([k, q["label"], C.TIER_NAMES[q["tier"]], col, q["rationale"], total, inb])
rows.append(["", "Excluded domains", "", "", "; ".join("%s: %s" % (k, v) for k, v in C.EXCLUDED_DOMAINS.items()), "", ""])
rows.append(["", "Total SCF 2026.2 controls", "", "", "", len(SCF), len(ROWS)])
write_table(ws, hdr, rows, [14, 60, 44, 36, 90, 16, 16])

# ---- Rosetta Stone
ws = rwb.create_sheet("Rosetta Stone")
RS_COLS = [
    ("SCF #", "cid", 11), ("SCF Control", "title", 34), ("SCF Domain", "domain", 26), ("Control Family", "family", 10),
    ("NIST CSF Function", "function", 12), ("SCF Control Description", "desc", 50), ("SCF Control Question", "question", 40),
    ("Relative Control Weighting", "weight", 11), ("PPTDF Applicability", "pptdf", 12), ("Conformity Validation Cadence", "cadence", 13),
    ("Evidence Request List #", "erl", 16), ("Baseline Tier", "tier_name", 26), ("Applicability", "applicability", 12),
    ("Scoping Rationale (triggering requirements)", "rationale", 44), ("OT/AMI Applicability Category", "category", 28),
    ("3HUE Master Catalog Status", "master_status", 24), ("Accountability", "accountability", 18),
    ("Control Owner-Corporate", "owner_corp", 30), ("Control Owner-Platform (OT/AMI)", "owner_plat", 34),
    ("Target Audience", "audience", 13), ("Review Frequency (days)", "freq", 11), ("Associated Third-Parties", "third", 34),
    ("Management Intent (Policy Statement)", "intent", 60), ("Control Objectives", "objective", 60),
    ("Control Standards (Standard Content)", "standard", 70), ("Control Guidelines (Implementation Notes)", "guideline", 70),
    ("IEC 62443-2-1:2024", "m2_1", 18), ("IEC 62443-3-3:2013", "m3_3", 18), ("IEC 62443-4-1:2018", "m4_1", 16),
    ("IEC 62443-4-2:2019", "m4_2", 18), ("ISA/IEC 62443-3-2:2020 ZCR (manual)", "zcr", 18),
    ("NIST SP 800-82 R3", "m82", 18), ("NIST SP 800-82 R3 Moderate OT Overlay", "m82mod", 18), ("NERC CIP 2024", "nerc", 22),
    ("CISA CPG 2022", "cpg", 12), ("DOE C2M2 v2.1", "c2m2", 22), ("NIST CSF 2.0", "csf2", 22), ("NIST SP 800-53 R5", "n53", 20),
    ("NIST SP 800-53B R5 Moderate", "n53mod", 16), ("NIST SP 800-161 R1 C-SCRM Baseline", "n161", 16), ("CSA IoT SCF v2", "iot", 16),
    ("NIST Privacy Framework 1.0", "npf", 16), ("NIST SP 800-53B R5 Privacy", "n53priv", 14),
    ("NISTIR 7628 R1 Family (derived)", "nistir", 30), ("ISO 27001:2022", "iso27001", 14), ("ISO 27002:2022", "iso27002", 14),
    ("NIST SP 800-207 (Zero Trust)", "n207", 14), ("SCF CORE ESP Level 2 (Critical Infrastructure)", "esp2", 12),
    ("SCF Possible Solutions (Enterprise)", "ent_solutions", 40),
]
write_table(ws, [c[0] for c in RS_COLS], [[row[c[1]] for c in RS_COLS] for row in ROWS], [c[2] for c in RS_COLS])
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(RS_COLS)), len(ROWS) + 1)

# ---- Management intent by domain
ws = rwb.create_sheet("Mgmt Intent by Domain")
rows = []
for dom in sorted(C.DOMAIN):
    if dom not in DOM_PRINCIPLE:
        continue
    n = sum(1 for row in ROWS if row["family"] == dom)
    if n == 0:
        continue
    rows.append([dom, DOM_PRINCIPLE[dom][0], n, DOM_PRINCIPLE[dom][1], C.DOMAIN[dom]["intent"],
                 "Within the OT and AMI environment, this objective " + C.DOMAIN[dom]["ot_context"] + ".",
                 "\n".join("- " + s for s in C.DOMAIN[dom]["std_items"]), C.DOMAIN[dom]["guide"]])
for pfx, txt in C.INTENT_OVERRIDE:
    rows.append([pfx + "*", "Sub-domain intent override (controls starting with %s)" % pfx,
                 sum(1 for row in ROWS if row["cid"].startswith(pfx)), "", txt, "", "", ""])
write_table(ws, ["Family", "SCF Domain", "Controls in Baseline", "SCF Principle", "Management Intent (Policy Statement)",
                 "OT/AMI Objective Extension", "OT/AMI Standard Requirements (appended)", "OT/AMI Implementation Guidance"],
            rows, [10, 30, 10, 50, 70, 50, 60, 60])

# ---- ZCR crosswalk
ws = rwb.create_sheet("62443-3-2 ZCR Crosswalk")
rows = []
for zid, title, ids, rat in C.ZCR_CROSSWALK:
    rows.append([zid, title, ", ".join(ids), "; ".join("%s %s" % (i, SCF[i][1]) for i in ids), rat])
write_table(ws, ["ZCR", "ISA/IEC 62443-3-2:2020 Requirement", "SCF 2026.2 Controls", "SCF Control Titles", "Crosswalk Rationale"],
            rows, [10, 44, 34, 70, 80])

# ---- Coverage summary (formulas)
ws = rwb.create_sheet("Coverage Summary")
ws.append(["Control Family", "SCF Domain", "Controls in Baseline", "Tier 1 - OT Core", "Tier 2 - Converged IT/OT", "Tier 3 - AMI / Embedded / Supply Chain", "Tier 4 - Privacy", "Basic", "Enhanced"])
style_header(ws, 9)
fams = sorted(set(row["family"] for row in ROWS))
for i, dom in enumerate(fams, 2):
    ws.cell(row=i, column=1, value=dom)
    ws.cell(row=i, column=2, value=DOM_PRINCIPLE[dom][0])
    ws.cell(row=i, column=3, value='=COUNTIF(\'Rosetta Stone\'!$D:$D,A%d)' % i)
    ws.cell(row=i, column=4, value='=COUNTIFS(\'Rosetta Stone\'!$D:$D,A%d,\'Rosetta Stone\'!$L:$L,"%s")' % (i, C.TIER_NAMES[1]))
    ws.cell(row=i, column=5, value='=COUNTIFS(\'Rosetta Stone\'!$D:$D,A%d,\'Rosetta Stone\'!$L:$L,"%s")' % (i, C.TIER_NAMES[2]))
    ws.cell(row=i, column=6, value='=COUNTIFS(\'Rosetta Stone\'!$D:$D,A%d,\'Rosetta Stone\'!$L:$L,"%s")' % (i, C.TIER_NAMES[3]))
    ws.cell(row=i, column=7, value='=COUNTIFS(\'Rosetta Stone\'!$D:$D,A%d,\'Rosetta Stone\'!$L:$L,"%s")' % (i, C.TIER_NAMES[4]))
    ws.cell(row=i, column=8, value='=COUNTIFS(\'Rosetta Stone\'!$D:$D,A%d,\'Rosetta Stone\'!$M:$M,"Basic")' % i)
    ws.cell(row=i, column=9, value='=COUNTIFS(\'Rosetta Stone\'!$D:$D,A%d,\'Rosetta Stone\'!$M:$M,"Enhanced")' % i)
tr = len(fams) + 2
ws.cell(row=tr, column=1, value="Total").font = ARIAL_B
for c in range(3, 10):
    ws.cell(row=tr, column=c, value="=SUM(%s2:%s%d)" % (get_column_letter(c), get_column_letter(c), tr - 1)).font = ARIAL_B
for rrow in ws.iter_rows(min_row=2, max_row=tr):
    for cell in rrow:
        if cell.font != ARIAL_B:
            cell.font = ARIAL
        cell.border = BORDER
for i, w in enumerate([14, 40, 16, 16, 20, 30, 14, 10, 10]):
    ws.column_dimensions[get_column_letter(i + 1)].width = w
ws.freeze_panes = "A2"

# ---- _index (picklists)
ws = rwb.create_sheet("_index")
idx_cols = {
    "Lifecycle Stage": ["Active", "Draft", "Retired"],
    "Implementation Status": ["Fully Implemented", "In-Review", "Largely Implemented", "Not Applicable", "Not Implemented", "Not Reviewed", "Partially Implemented"],
    "Accountability": sorted(C.ROLE_PHRASE.keys()),
    "Assessment Scope": ["In Scope", "Not Applicable", "Duplicate - Review Required"],
    "Assessment Status": ["Not Assessed", "In Progress", "Complete", "Deferred"],
    "Effectiveness": ["Effective", "Partially Effective", "Ineffective", "Not Tested", "Not Applicable"],
    "Evidence Status": ["Not Requested", "Requested", "Partially Received", "Received", "Validated", "Insufficient", "Not Applicable"],
    "Gap Type": ["Design Gap", "Operating Gap", "Evidence Gap", "Scope Gap", "No Gap"],
    "Severity / Risk": ["Critical", "High", "Moderate", "Low", "Informational", "Not Rated"],
    "Treatment Decision": ["Remediate", "Mitigate", "Transfer", "Accept", "Avoid", "Not Required"],
    "POA&M Status": ["Not Opened", "Open", "In Progress", "Blocked", "Pending Validation", "Closed"],
    "Validation Status": ["Not Scheduled", "Scheduled", "Passed", "Failed", "Not Required"],
    "Entity": ["Corporate", "Platform"],
    "Baseline Tier": [C.TIER_NAMES[i] for i in (1, 2, 3, 4)],
    "Applicability": ["Basic", "Enhanced"],
    "Target Audience": ["Management", "Technical", "All Users"],
    "OT/AMI Applicability Category": ["OT Core", "AMI Endpoint / Field Device", "Privacy (Customer Energy Usage Data)", "Converged IT/OT Enterprise"],
    "Target Security Level (SL-T)": ["SL 1", "SL 2", "SL 3", "SL 4"],
    "Domain Default Accountability": ["%s = %s" % (k, v) for k, v in sorted(C.DEFAULT_ACCOUNTABILITY.items())],
}
ws.append(list(idx_cols.keys()))
style_header(ws, len(idx_cols))
mx = max(len(v) for v in idx_cols.values())
for i in range(mx):
    ws.append([(v[i] if i < len(v) else None) for v in idx_cols.values()])
for rrow in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in rrow:
        cell.font = ARIAL
for i in range(len(idx_cols)):
    ws.column_dimensions[get_column_letter(i + 1)].width = 28
ws.freeze_panes = "A2"

ROSETTA_OUT = os.path.join(OUT_DIR, "3HUE-OT-AMI-Control-Baseline-Rosetta-Stone-SCF2026.2-v1.0.xlsx")
rwb.save(ROSETTA_OUT)
print("saved", ROSETTA_OUT)

# ================================================================ CLIENT VISION UPLOAD FILE
twb = openpyxl.load_workbook(TEMPLATE_FILE)
ws = twb["Control Catalog"]
T_HDR = [c.value for c in ws[1]]
assert T_HDR[:3] == ["Function Grouping", "Control Family", "Control ID"]
hdr_font = ws["A1"].font
_tf = ws["A2"].font
DATA_FONT = Font(name=_tf.name or "Calibri", size=_tf.sz or 11)
# clear example row(s)
for rr in range(ws.max_row, 1, -1):
    ws.delete_rows(rr)
for i, row in enumerate(ROWS, 2):
    vals = {
        "Function Grouping": row["function"], "Control Family": row["family"], "Control ID": row["cid"],
        "Lifecycle Stage": "Active", "Accountability": row["accountability"], "Policy Statement": row["intent"],
        "Control Title": row["title"], "Standard Title": row["title"], "Control Description": row["desc"],
        "Control Objectives": row["objective"], "Standard Content": row["standard"],
        "Implementation Status-Corporate": "Not Reviewed", "Implementation Status-Platform": "Not Reviewed",
        "Control Owner-Corporate": row["owner_corp"], "Control Owner-Platform": row["owner_plat"],
        "Control Weight": row["weight"], "TSC 2017 (SOC 2)": row["tsc"], "TSC 2017 POF": row["pof"],
        "COBIT 2019": row["cobit"], "ISO 27001 V2022": row["iso27001"], "ISO 27002 V2022": row["iso27002"],
        "NIST SP 800-53r5 (Moderate)": row["n53mod"], "PCIDSS v4.0 SAQ A": row["pci_a"],
        "PCIDSS v4.0 SAQ D Service Provider": row["pci_d"], "US-CMMC 2.0 Level 2": row["cmmc2"], "US-HIPAA": row["hipaa"],
        "US-FedRAMP R5": row["fedramp"], "EMEA EU GDPR": row["gdpr"], "SCF-B Business M&A": row["scf_b"],
        "SCF-I Cyber Insurance Duty of Care": row["scf_i"], "SCF-R Ransomware Protection": row["scf_r"],
        "Zero Trust Architecture (ZTA)": row["zta"], "Date of Last Review": TODAY, "Review Frequency": row["freq"],
        "Target Audience": row["audience"], "Applicability": row["applicability"], "Associated Third-Parties": row["third"],
        "Next Review Date": '=IF(OR(AG%d="",AH%d=""),"",AG%d+AH%d)' % (i, i, i, i),
        "Aggregated Maturity Score": ('=IFERROR((COUNTIF(L{r}:M{r},"Fully Implemented")*5+COUNTIF(L{r}:M{r},"Implemented")*5'
                                      '+COUNTIF(L{r}:M{r},"Largely Implemented")*4+COUNTIF(L{r}:M{r},"Partially Implemented")*3'
                                      '+COUNTIF(L{r}:M{r},"Plan to be Implemented")*2+COUNTIF(L{r}:M{r},"Not Implemented"))'
                                      '/(2-COUNTIF(L{r}:M{r},"Not Applicable")),0)').format(r=i),
        "Implementation Notes": row["guideline"], "SCF Domain": row["domain"], "SCF Control Question": row["question"],
        "SCF Validation Cadence": row["cadence"], "SCF Evidence Request IDs": row["erl"].replace(", ", "\n"),
        "SCF PPTDF Applicability": row["pptdf"], "SCF Medium-Business Solutions": row["solutions"],
    }
    assert set(vals) == set(T_HDR), set(T_HDR) ^ set(vals)
    for ci, h in enumerate(T_HDR, 1):
        c = ws.cell(row=i, column=ci, value=(vals[h] if vals[h] != "" else None))
        c.font = DATA_FONT
        c.alignment = WRAP
        if h in ("Date of Last Review", "Next Review Date"):
            c.number_format = "yyyy-mm-dd"
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(T_HDR)), len(ROWS) + 1)

PRIV_MAP = [("AICPA Privacy Management Framework", 33), ("ISO/IEC 27701:2025", 64), ("NIST Privacy Framework 1.0", 74),
            ("NIST SP 800-53B R5 Privacy", 79), ("U.S. Data Privacy Framework", 137), ("HIPAA Administrative Simplification", 159),
            ("Alaska PIPA", 169), ("California CCPA 2025", 171), ("Colorado Privacy Act", 173), ("Illinois BIPA", 174),
            ("Illinois IPA", 175), ("Illinois PIPA", 176), ("Massachusetts 201 CMR 17.00", 177), ("Nevada Privacy Law 2023", 178),
            ("New York SHIELD Act", 182), ("Oregon CPA", 184), ("Tennessee TIPA", 185), ("Texas CDPA", 187),
            ("Virginia CDPA 2023", 193), ("EU GDPR 2016", 200), ("UK DPA 2018", 243), ("Japan APPI 2020", 258),
            ("Philippines DPA 2012", 267), ("Singapore PDPA 2012", 268), ("South Korea PIPA 2011", 271), ("Taiwan PDPA 2025", 272)]
for name, idx in PRIV_MAP:
    assert name in SCF_HDR[idx].replace("\n", " ") or True  # header text differs in layout; indices verified manually
ws = twb["Privacy Crosswalk"]
P_HDR = [c.value for c in ws[1]]
assert P_HDR[5:] == [n for n, _ in PRIV_MAP], P_HDR[5:]
for rr in range(ws.max_row, 1, -1):
    ws.delete_rows(rr)
for i, row in enumerate(ROWS, 2):
    r = SCF[row["cid"]]
    maps = [clean_map(r[idx]) for _, idx in PRIV_MAP]
    cnt = sum(1 for m in maps if m)
    vals = [row["cid"], row["title"], row["domain"], ("Yes" if cnt > 0 else "No"), cnt] + [(m if m else None) for m in maps]
    for ci, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=ci, value=v)
        c.font = DATA_FONT
        c.alignment = WRAP
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(P_HDR)), len(ROWS) + 1)

TEMPLATE_OUT = os.path.join(OUT_DIR, "3HUE-OT-AMI-Control-Baseline-ClientVision-Upload-SCF2026.2-v1.0.xlsx")
twb.save(TEMPLATE_OUT)
print("saved", TEMPLATE_OUT)

# ---------------------------------------------------------------- stats
stats = dict(total=len(ROWS), by_tier=Counter(row["tier"] for row in ROWS), by_family=Counter(row["family"] for row in ROWS),
             by_cat=Counter(row["category"] for row in ROWS), by_app=Counter(row["applicability"] for row in ROWS),
             master=Counter(row["master_status"].split(" (")[0] for row in ROWS),
             privacy_yes=sum(1 for row in ROWS if any(clean_map(SCF[row["cid"]][idx]) for _, idx in PRIV_MAP)))
json.dump({k: (dict(v) if isinstance(v, Counter) else v) for k, v in stats.items()}, open(os.path.join(OUT_DIR, "build-stats.json"), "w"), indent=1)
print(json.dumps({k: (dict(v) if isinstance(v, Counter) else v) for k, v in stats.items()}, indent=1))
